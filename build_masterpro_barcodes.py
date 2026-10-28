from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

INPUT = Path(r"C:\Users\<user>\Downloads\MasterPro_UPC.xlsx")
OUTPUT_DIR = Path(r"C:\Users\<user>\Desktop\Adattauto\output")
BARCODE_DIR = OUTPUT_DIR / "masterpro_barcodes"
OUTPUT = OUTPUT_DIR / "MasterPro_UPC_with_barcodes.xlsx"

L_CODES = {
    "0": "0001101", "1": "0011001", "2": "0010011", "3": "0111101", "4": "0100011",
    "5": "0110001", "6": "0101111", "7": "0111011", "8": "0110111", "9": "0001011",
}
R_CODES = {
    "0": "1110010", "1": "1100110", "2": "1101100", "3": "1000010", "4": "1011100",
    "5": "1001110", "6": "1010000", "7": "1000100", "8": "1001000", "9": "1110100",
}


def gtin14_to_upc12(value):
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    text = "".join(ch for ch in text if ch.isdigit())
    if len(text) == 14 and text.startswith("00"):
        return text[2:]
    if len(text) == 12:
        return text
    raise ValueError(f"GTIN must be 14 digits starting with 00 or 12 UPC digits: {value}")


def valid_upc_check_digit(upc):
    if len(upc) != 12 or not upc.isdigit():
        return False
    digits = [int(ch) for ch in upc]
    total = sum(digits[i] * 3 for i in range(0, 11, 2)) + sum(digits[i] for i in range(1, 11, 2))
    check = (10 - (total % 10)) % 10
    return check == digits[-1]


def make_upc_pattern(upc):
    return (
        "101"
        + "".join(L_CODES[d] for d in upc[:6])
        + "01010"
        + "".join(R_CODES[d] for d in upc[6:])
        + "101"
    )


def make_barcode(upc, output_path):
    module = 3
    quiet = 10 * module
    bar_height = 82
    text_height = 22
    pattern = make_upc_pattern(upc)
    width = quiet * 2 + len(pattern) * module
    height = bar_height + text_height + 12
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 18)
        small_font = ImageFont.truetype("arial.ttf", 15)
    except Exception:
        font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    x = quiet
    for idx, bit in enumerate(pattern):
        is_guard = idx < 3 or 45 <= idx < 50 or idx >= 92
        h = bar_height + 8 if is_guard else bar_height
        if bit == "1":
            draw.rectangle([x, 6, x + module - 1, h], fill="black")
        x += module

    # Human-readable UPC-A layout: first digit left, five + five center, last digit right.
    draw.text((quiet - 18, bar_height - 2), upc[0], fill="black", font=small_font)
    draw.text((quiet + 16 * module, bar_height + 5), upc[1:6], fill="black", font=font)
    draw.text((quiet + 50 * module + 10 * module, bar_height + 5), upc[6:11], fill="black", font=font)
    draw.text((quiet + 95 * module + 3, bar_height - 2), upc[11], fill="black", font=small_font)
    image.save(output_path)


def style_header(ws, end_col):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for cell in ws[1][:end_col]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    BARCODE_DIR.mkdir(parents=True, exist_ok=True)

    source = load_workbook(INPUT, data_only=True)
    products = source["Products"]
    headers = [cell.value for cell in products[1]]
    col = {name: idx + 1 for idx, name in enumerate(headers)}

    required = ["GTIN", "SKU", "Description", "BrandName"]
    missing = [name for name in required if name not in col]
    if missing:
        raise RuntimeError(f"Missing required source columns: {missing}")

    out = Workbook()
    ws = out.active
    ws.title = "UPC Barcodes"
    source_ws = out.create_sheet("Products Source")

    ws_headers = ["SKU", "Description", "BrandName", "GTIN14", "UPC12", "CheckDigitOK", "Barcode"]
    ws.append(ws_headers)

    rows = []
    for row_idx in range(2, products.max_row + 1):
        gtin = products.cell(row_idx, col["GTIN"]).value
        if not gtin:
            continue
        sku = products.cell(row_idx, col["SKU"]).value
        desc = products.cell(row_idx, col["Description"]).value
        brand = products.cell(row_idx, col["BrandName"]).value
        upc = gtin14_to_upc12(gtin)
        ok = "Y" if valid_upc_check_digit(upc) else "N"
        rows.append((sku, desc, brand, str(gtin), upc, ok))

    for idx, row in enumerate(rows, start=2):
        sku, desc, brand, gtin, upc, ok = row
        ws.append([sku, desc, brand, gtin, upc, ok, ""])
        barcode_path = BARCODE_DIR / f"{idx-1:03d}_{upc}.png"
        make_barcode(upc, barcode_path)
        img = XLImage(str(barcode_path))
        img.width = 260
        img.height = 115
        ws.add_image(img, f"G{idx}")
        ws.row_dimensions[idx].height = 92

    source_ws.append(headers)
    for source_row in products.iter_rows(min_row=2, values_only=True):
        source_ws.append(list(source_row))

    for tab in [ws, source_ws]:
        tab.freeze_panes = "A2"
        tab.sheet_view.showGridLines = False
        end_col = tab.max_column
        style_header(tab, end_col)
        thin = Side(style="thin", color="D9E2F3")
        for row in tab.iter_rows(min_row=1, max_row=tab.max_row, max_col=end_col):
            for cell in row:
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                cell.font = cell.font.copy(name="Arial", size=10)
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = {
        "A": 32, "B": 28, "C": 16, "D": 18, "E": 16, "F": 14, "G": 39,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for col_idx in range(1, source_ws.max_column + 1):
        source_ws.column_dimensions[get_column_letter(col_idx)].width = 18
    for col_letter in ["D", "E"]:
        for cell in ws[col_letter][1:]:
            cell.number_format = "@"

    out.save(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"rows={len(rows)}")
    print(f"invalid_check_digits={sum(1 for row in rows if row[-1] != 'Y')}")


if __name__ == "__main__":
    main()
